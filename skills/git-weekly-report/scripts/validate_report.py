#!/usr/bin/env python3
"""Validate a weekly report before delivery.

Checks three layers:

1. report.json: required fields, status/deviation consistency, evidence.
2. coverage: every project with commits is referenced by at least one item.
3. Excel: title, headers, section merges, row heights, print setup, and an
   estimate that the content still fits on one A4 landscape page.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - depends on the caller's environment
    print(
        "[error] openpyxl is required. Run with: uv run --with openpyxl python ...",
        file=sys.stderr,
    )
    raise SystemExit(2)

HEADERS = ["大类", "名称（只列关键）", "目标", "进展", "完成状态", "偏差说明"]
COLUMN_WIDTHS = {"A": 8, "B": 24, "C": 34, "D": 68, "E": 10, "F": 22}
STATUS_VALUES = {"完成", "进行中", "部分完成", "有偏差"}
EMPTY_DEVIATIONS = {"", "无", "无。", "-", "none", "n/a"}

DEFAULT_CONFIG = Path.home() / ".config" / "git-weekly-report" / "config.json"

A4_LANDSCAPE_WIDTH_IN = 11.69
A4_LANDSCAPE_HEIGHT_IN = 8.27
MARGIN_LEFT_IN = 0.3
MARGIN_RIGHT_IN = 0.3
MARGIN_TOP_IN = 0.4
MARGIN_BOTTOM_IN = 0.4
EXCEL_WIDTH_UNIT_PX = 7
EXCEL_COLUMN_PADDING_PX = 5
EXCEL_DPI = 96
POINTS_PER_INCH = 72


class Validation:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warning(self, message: str) -> None:
        self.warnings.append(message)

    def ok(self) -> bool:
        return not self.errors


def load_json(path: Path, label: str) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"{label} not found: {path}")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON in {label}: {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError(f"{label} must be a JSON object: {path}")
    return data


def load_config(path: Path, *, required: bool) -> dict[str, Any]:
    if not path.exists():
        if required:
            raise FileNotFoundError(f"config not found: {path}")
        return {}
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON in config: {path}: {exc}") from exc
    if not isinstance(config, dict):
        raise ValueError(f"config must be a JSON object: {path}")
    return config


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def vertical_label(label: str) -> str:
    compact = label.strip()
    if 1 < len(compact) <= 6 and all(
        "\u4e00" <= char <= "\u9fff" or char in " ·/-" for char in compact
    ):
        return "\n".join(compact)
    return compact


def check_report_structure(
    report: dict[str, Any],
    commits: dict[str, Any],
    result: Validation,
) -> tuple[set[str], int, dict[str, int]]:
    period = commits.get("period", {})
    report_period = text(report.get("period"))
    if not report_period:
        result.error("report.period is empty")
    elif report_period != text(period.get("label")):
        result.error(
            f"report.period {report_period!r} does not match commits.period.label "
            f"{text(period.get('label'))!r}"
        )

    report_week = text(report.get("week"))
    commits_week = text(period.get("week"))
    if report_week and commits_week and report_week != commits_week:
        result.error(
            f"report.week {report_week!r} does not match commits.period.week {commits_week!r}"
        )

    if period.get("multi_week"):
        result.error("the selected range spans more than one ISO week; use a single week")

    sections = report.get("sections")
    if not isinstance(sections, list) or not sections:
        result.error("report.sections must be a non-empty list")
        return set(), 0, {}

    by_project = commits.get("stats", {}).get("by_project", {})
    if not isinstance(by_project, dict):
        by_project = {}

    covered_projects: set[str] = set()
    section_counts: dict[str, int] = {}
    total_items = 0

    for section_index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            result.error(f"sections[{section_index}] must be an object")
            continue
        label = text(section.get("label"))
        if not label:
            result.error(f"sections[{section_index}].label is empty")
        items = section.get("items")
        if not isinstance(items, list) or not items:
            result.error(f"sections[{section_index}].items must be a non-empty list")
            continue
        section_counts[label] = section_counts.get(label, 0) + len(items)
        total_items += len(items)

        for item_index, item in enumerate(items, start=1):
            prefix = f"sections[{section_index}].items[{item_index}]"
            if not isinstance(item, dict):
                result.error(f"{prefix} must be an object")
                continue

            for field in ("name", "goal", "progress", "status", "deviation"):
                if not text(item.get(field)):
                    result.error(f"{prefix}.{field} is empty")

            status = text(item.get("status"))
            if status and status not in STATUS_VALUES:
                result.error(
                    f"{prefix}.status must be one of {sorted(STATUS_VALUES)}, got {status!r}"
                )
            deviation = text(item.get("deviation"))
            if status and status != "完成" and deviation.lower() in EMPTY_DEVIATIONS:
                result.error(
                    f"{prefix}.deviation must explain why status is {status!r}, not {deviation!r}"
                )

            evidence = item.get("evidence")
            if not isinstance(evidence, dict):
                result.error(f"{prefix}.evidence must be an object")
                continue

            projects = evidence.get("projects")
            if not isinstance(projects, list) or not projects:
                result.error(f"{prefix}.evidence.projects must be a non-empty list")
                projects = []
            normalized_projects = [text(project) for project in projects if text(project)]
            if len(normalized_projects) != len(projects):
                result.error(f"{prefix}.evidence.projects contains an empty project name")

            for project in normalized_projects:
                if project not in by_project:
                    result.error(
                        f"{prefix}.evidence.projects references {project!r}, "
                        "which has no commits in commits.json"
                    )
                else:
                    covered_projects.add(project)

            commit_count = evidence.get("commit_count")
            if not isinstance(commit_count, int) or isinstance(commit_count, bool):
                result.error(f"{prefix}.evidence.commit_count must be an integer")
            elif commit_count <= 0:
                result.error(f"{prefix}.evidence.commit_count must be greater than 0")
            else:
                actual = sum(int(by_project.get(project, 0)) for project in normalized_projects)
                if commit_count > actual:
                    result.error(
                        f"{prefix}.evidence.commit_count {commit_count} exceeds "
                        f"the actual {actual} commit(s) in {normalized_projects}"
                    )

            samples = evidence.get("sample_commits")
            if not isinstance(samples, list) or not samples:
                result.error(f"{prefix}.evidence.sample_commits must be a non-empty list")
            else:
                min_samples = (
                    1
                    if isinstance(commit_count, int) and not isinstance(commit_count, bool) and commit_count == 1
                    else 2
                )
                if not min_samples <= len(samples) <= 5:
                    result.warning(
                        f"{prefix}.evidence.sample_commits should have "
                        f"{min_samples}-5 items, got {len(samples)}"
                    )
                if not all(text(sample) for sample in samples):
                    result.error(f"{prefix}.evidence.sample_commits contains an empty value")

            repos = evidence.get("repos")
            if repos is not None:
                if not isinstance(repos, list) or not all(text(repo) for repo in repos):
                    result.error(f"{prefix}.evidence.repos must be a list of repo paths")
                else:
                    by_repo = commits.get("stats", {}).get("by_repo", {})
                    if isinstance(by_repo, dict):
                        for repo in repos:
                            repo_name = text(repo)
                            if repo_name and repo_name not in by_repo:
                                result.warning(
                                    f"{prefix}.evidence.repos references {repo_name!r}, "
                                    "which has no commits in commits.json"
                                )

    return covered_projects, total_items, section_counts


def check_coverage(
    commits: dict[str, Any],
    covered_projects: set[str],
    result: Validation,
) -> None:
    unmapped = commits.get("unmapped", [])
    if unmapped:
        repos = ", ".join(
            f"{item.get('repo')}({item.get('commit_count')})"
            for item in unmapped
            if isinstance(item, dict)
        )
        result.error(f"unmapped repos must be confirmed before generating the report: {repos}")

    stats = commits.get("stats", {})
    by_project = stats.get("by_project", {})
    if isinstance(by_project, dict):
        missing = sorted(
            project
            for project, count in by_project.items()
            if int(count) > 0 and project not in covered_projects
        )
        if missing:
            result.error(
                "projects with commits are not covered by any report item: "
                + ", ".join(missing)
            )

    if stats.get("errors"):
        for error in stats["errors"]:
            result.warning(f"extract warning: {error}")

    if stats.get("missing_repos"):
        result.warning(
            "project map contains repos that do not exist in the workspace: "
            + ", ".join(stats["missing_repos"])
        )

    total = stats.get("total")
    commit_list = commits.get("commits", [])
    if isinstance(total, int) and isinstance(commit_list, list) and total != len(commit_list):
        result.warning(
            f"stats.total {total} does not match len(commits) {len(commit_list)}"
        )


def check_recommended_counts(
    section_counts: dict[str, int],
    total_items: int,
    result: Validation,
) -> None:
    key_count = section_counts.get("重点项目")
    daily_count = section_counts.get("日常工作")
    if key_count is not None and not 3 <= key_count <= 5:
        result.warning(f"重点项目 should have 3-5 items, got {key_count}")
    if daily_count is not None and not 2 <= daily_count <= 4:
        result.warning(f"日常工作 should have 2-4 items, got {daily_count}")
    if total_items > 8:
        result.warning(f"total items should be <= 8 to fit one page, got {total_items}")


def check_samples(
    report: dict[str, Any],
    commits: dict[str, Any],
    result: Validation,
) -> None:
    by_project: dict[str, set[str]] = {}
    for commit in commits.get("commits", []):
        if not isinstance(commit, dict):
            continue
        project = text(commit.get("project"))
        if not project:
            continue
        by_project.setdefault(project, set()).add(text(commit.get("norm")))
        by_project[project].add(text(commit.get("subject")))

    for section in report.get("sections", []):
        if not isinstance(section, dict):
            continue
        for item in section.get("items", []):
            if not isinstance(item, dict):
                continue
            evidence = item.get("evidence", {})
            projects = evidence.get("projects", []) if isinstance(evidence, dict) else []
            samples = evidence.get("sample_commits", []) if isinstance(evidence, dict) else []
            corpus: set[str] = set()
            for project in projects:
                corpus.update(by_project.get(text(project), set()))
            if not corpus or not isinstance(samples, list):
                continue
            matched = 0
            for sample in samples:
                sample_text = text(sample)
                if not sample_text:
                    continue
                if any(sample_text in entry or entry in sample_text for entry in corpus if entry):
                    matched += 1
            if matched == 0:
                result.warning(
                    f"sample_commits for {text(item.get('name'))!r} do not match any commit title"
                )


def check_excel(
    xlsx_path: Path,
    report: dict[str, Any],
    result: Validation,
    *,
    sheet_name: str,
) -> None:
    if not xlsx_path.exists():
        result.error(f"xlsx not found: {xlsx_path}")
        return

    workbook = load_workbook(xlsx_path)
    if sheet_name not in workbook.sheetnames:
        result.error(
            f"worksheet {sheet_name!r} not found; available: {', '.join(workbook.sheetnames)}"
        )
        return
    worksheet = workbook[sheet_name]

    expected_title = f"周报周期：{text(report.get('period'))}"
    actual_title = text(worksheet["A1"].value)
    if actual_title != expected_title:
        result.error(f"A1 title {actual_title!r} does not match {expected_title!r}")

    merged = {str(item) for item in worksheet.merged_cells.ranges}
    if "A1:F1" not in merged:
        result.error("title row A1:F1 is not merged")

    headers = [text(worksheet.cell(row=2, column=column).value) for column in range(1, 7)]
    if headers != HEADERS:
        result.error(f"header row does not match: {headers!r}")

    row = 3
    last_expected_row = 2
    for section_index, section in enumerate(report.get("sections", [])):
        if not isinstance(section, dict):
            continue
        items = section.get("items", [])
        if not isinstance(items, list) or not items:
            continue
        if section_index > 0:
            separator_height = worksheet.row_dimensions[row].height
            if separator_height is None or not 4 <= float(separator_height) <= 12:
                result.warning(
                    f"separator row {row} should have height around 8, got {separator_height!r}"
                )
            row += 1

        start_row = row
        end_row = start_row + len(items) - 1
        last_expected_row = end_row

        if end_row > start_row:
            expected_merge = f"A{start_row}:A{end_row}"
            if expected_merge not in merged:
                result.error(f"section A column is not merged: {expected_merge}")

        expected_label = vertical_label(text(section.get("label")))
        actual_label = text(worksheet.cell(row=start_row, column=1).value)
        if actual_label != expected_label:
            result.error(
                f"A{start_row} section label {actual_label!r} does not match {expected_label!r}"
            )

        for offset, item in enumerate(items):
            current = start_row + offset
            if not isinstance(item, dict):
                continue
            expected = {
                2: text(item.get("name")),
                3: text(item.get("goal")),
                4: text(item.get("progress")),
                5: text(item.get("status")),
                6: text(item.get("deviation")),
            }
            for column, expected_value in expected.items():
                actual_value = text(worksheet.cell(row=current, column=column).value)
                if actual_value != expected_value:
                    result.error(
                        f"{worksheet.cell(row=current, column=column).coordinate} "
                        f"does not match report item {text(item.get('name'))!r}"
                    )
            height = worksheet.row_dimensions[current].height
            if height is None:
                result.warning(f"row {current} has no explicit height")
            elif float(height) < 30:
                result.warning(f"row {current} height {height} is below 30")
            elif float(height) > 409:
                result.error(f"row {current} height {height} exceeds Excel's 409 limit")

        row = end_row + 1

    if worksheet.max_row > last_expected_row:
        result.warning(
            f"worksheet has {worksheet.max_row} rows, expected {last_expected_row}"
        )
    elif worksheet.max_row < last_expected_row:
        result.error(
            f"worksheet has only {worksheet.max_row} rows, expected {last_expected_row}"
        )

    if worksheet.page_setup.orientation != "landscape":
        result.error(
            f"page orientation must be landscape, got {worksheet.page_setup.orientation!r}"
        )
    paper_size = worksheet.page_setup.paperSize
    if paper_size is None or int(paper_size) != int(worksheet.PAPERSIZE_A4):
        result.error(
            f"paper size must be A4 ({worksheet.PAPERSIZE_A4}), "
            f"got {paper_size!r}"
        )
    if worksheet.page_setup.fitToWidth != 1:
        result.error(
            f"fitToWidth must be 1, got {worksheet.page_setup.fitToWidth!r}"
        )
    if worksheet.page_setup.fitToHeight != 0:
        result.error(
            f"fitToHeight must be 0, got {worksheet.page_setup.fitToHeight!r}"
        )
    page_setup = worksheet.sheet_properties.pageSetUpPr
    if page_setup is None or not page_setup.fitToPage:
        result.error("fitToPage must be enabled")
    if worksheet.sheet_view.showGridLines:
        result.error("grid lines must be hidden")
    print_titles = text(worksheet.print_title_rows).replace("$", "")
    if print_titles != "1:2":
        result.error(f"print_title_rows must be 1:2, got {worksheet.print_title_rows!r}")

    margins = worksheet.page_margins
    for name, actual, expected in (
        ("left", margins.left, MARGIN_LEFT_IN),
        ("right", margins.right, MARGIN_RIGHT_IN),
        ("top", margins.top, MARGIN_TOP_IN),
        ("bottom", margins.bottom, MARGIN_BOTTOM_IN),
    ):
        if actual is None or abs(float(actual) - expected) > 0.01:
            result.warning(f"page margin {name} should be {expected}, got {actual!r}")

    for column, expected_width in COLUMN_WIDTHS.items():
        actual_width = worksheet.column_dimensions[column].width
        if actual_width is None or abs(float(actual_width) - expected_width) > 0.1:
            result.warning(
                f"column {column} width should be {expected_width}, got {actual_width!r}"
            )

    check_single_page(worksheet, result)


def check_single_page(worksheet: Any, result: Validation) -> None:
    total_height = 0.0
    for row in range(1, worksheet.max_row + 1):
        height = worksheet.row_dimensions[row].height
        total_height += float(height) if height is not None else 15.0

    total_width_units = 0.0
    for column in COLUMN_WIDTHS:
        width = worksheet.column_dimensions[column].width
        total_width_units += float(width) if width is not None else COLUMN_WIDTHS[column]

    estimated_width_px = (
        total_width_units * EXCEL_WIDTH_UNIT_PX
        + len(COLUMN_WIDTHS) * EXCEL_COLUMN_PADDING_PX
    )
    estimated_width_in = estimated_width_px / EXCEL_DPI
    usable_width_in = A4_LANDSCAPE_WIDTH_IN - MARGIN_LEFT_IN - MARGIN_RIGHT_IN
    usable_height_pt = (
        A4_LANDSCAPE_HEIGHT_IN - MARGIN_TOP_IN - MARGIN_BOTTOM_IN
    ) * POINTS_PER_INCH

    scale = min(1.0, usable_width_in / estimated_width_in) if estimated_width_in else 1.0
    scaled_height_pt = total_height * scale

    print(
        "  page estimate: "
        f"raw={total_height:.1f}pt, width={estimated_width_in:.2f}in, "
        f"scale={scale:.2f}, scaled={scaled_height_pt:.1f}pt, "
        f"usable={usable_height_pt:.1f}pt"
    )

    if scaled_height_pt > usable_height_pt + 2:
        result.error(
            "estimated content exceeds one A4 landscape page "
            f"({scaled_height_pt:.1f}pt > {usable_height_pt:.1f}pt); "
            "shorten progress text or reduce row heights"
        )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Validate report.json, commits.json and the rendered weekly report.",
    )
    parser.add_argument("--report", required=True, help="report.json path")
    parser.add_argument("--commits", required=True, help="commits.json path")
    parser.add_argument("--xlsx", help="rendered .xlsx path; omit to skip Excel checks")
    parser.add_argument(
        "--config",
        help=f"local config JSON (default: {DEFAULT_CONFIG})",
    )
    parser.add_argument(
        "--sheet-name",
        default=None,
        help="worksheet name; defaults to config report.sheet_name or 上周周报",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    result = Validation()

    try:
        config_path = Path(args.config).expanduser() if args.config else DEFAULT_CONFIG
        config = load_config(config_path, required=bool(args.config))
        report = load_json(Path(args.report).expanduser(), "report")
        commits = load_json(Path(args.commits).expanduser(), "commits")
    except (FileNotFoundError, ValueError) as exc:
        print(f"[error] {exc}", file=sys.stderr)
        return 2

    sheet_name = args.sheet_name
    if not sheet_name:
        report_config = config.get("report", {})
        if isinstance(report_config, dict):
            sheet_name = str(report_config.get("sheet_name", "")).strip() or None
    sheet_name = sheet_name or "上周周报"

    covered_projects, total_items, section_counts = check_report_structure(
        report, commits, result
    )
    check_coverage(commits, covered_projects, result)
    check_recommended_counts(section_counts, total_items, result)
    check_samples(report, commits, result)

    if args.xlsx:
        check_excel(
            Path(args.xlsx).expanduser(),
            report,
            result,
            sheet_name=sheet_name,
        )
    else:
        result.warning("xlsx not provided; Excel layout was not checked")

    if result.errors:
        print(f"[FAIL] {len(result.errors)} error(s), {len(result.warnings)} warning(s)")
        for error in result.errors:
            print(f"  error: {error}")
        for warning in result.warnings:
            print(f"  warning: {warning}")
        return 1

    print(f"[PASS] {len(result.warnings)} warning(s)")
    for warning in result.warnings:
        print(f"  warning: {warning}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
