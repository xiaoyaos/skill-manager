#!/usr/bin/env python3
"""Render report.json into the fixed single-week Excel template.

This script only handles layout. It does not decide which commits matter or
how to phrase them. Edit report.json and re-render instead of editing the
generated .xlsx by hand.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

DEFAULT_CONFIG = Path.home() / ".config" / "git-weekly-report" / "config.json"

HEADERS = ["大类", "名称（只列关键）", "目标", "进展", "完成状态", "偏差说明"]
COLUMN_WIDTHS = {"A": 8, "B": 24, "C": 34, "D": 68, "E": 10, "F": 22}
STATUS_VALUES = {"完成", "进行中", "部分完成", "有偏差"}
FONT_NAME = "宋体"

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TITLE_FILL = PatternFill("solid", fgColor="D9E2F3")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="E2EFDA")
TITLE_COLOR = "1F4E79"
TEXT_COLOR = "1F1F1F"


class ReportError(RuntimeError):
    """Raised when report.json is missing required fields."""


def load_report(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise ReportError(f"report not found: {path}")
    try:
        report = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ReportError(f"invalid JSON in report: {path}: {exc}") from exc
    if not isinstance(report, dict):
        raise ReportError(f"report must be a JSON object: {path}")
    return report


def load_config(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ReportError(f"invalid JSON in config: {path}: {exc}") from exc
    if not isinstance(config, dict):
        raise ReportError(f"config must be a JSON object: {path}")
    return config


def report_filename(period: str) -> str:
    match = re.fullmatch(
        r"(\d{4})/(\d{2})/(\d{2})-(\d{4})/(\d{2})/(\d{2})",
        period.strip(),
    )
    if not match:
        raise ReportError(
            f"cannot derive output filename from period {period!r}; "
            "pass --output explicitly"
        )
    start = "".join(match.groups()[:3])
    end = "".join(match.groups()[3:])
    return f"周报_{start}-{end}.xlsx"


def require_text(value: Any, field: str) -> str:
    text = "" if value is None else str(value).strip()
    if not text:
        raise ReportError(f"required field is empty: {field}")
    return text


def validate_report(report: dict[str, Any]) -> None:
    require_text(report.get("period"), "period")
    sections = report.get("sections")
    if not isinstance(sections, list) or not sections:
        raise ReportError("'sections' must be a non-empty list")

    for section_index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            raise ReportError(f"sections[{section_index}] must be an object")
        require_text(section.get("label"), f"sections[{section_index}].label")
        items = section.get("items")
        if not isinstance(items, list) or not items:
            raise ReportError(f"sections[{section_index}].items must be a non-empty list")
        for item_index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                raise ReportError(
                    f"sections[{section_index}].items[{item_index}] must be an object"
                )
            prefix = f"sections[{section_index}].items[{item_index}]"
            require_text(item.get("name"), f"{prefix}.name")
            require_text(item.get("goal"), f"{prefix}.goal")
            require_text(item.get("progress"), f"{prefix}.progress")
            status = require_text(item.get("status"), f"{prefix}.status")
            if status not in STATUS_VALUES:
                raise ReportError(
                    f"{prefix}.status must be one of {sorted(STATUS_VALUES)}, got: {status!r}"
                )
            require_text(item.get("deviation"), f"{prefix}.deviation")


def text_units(text: str) -> int:
    return sum(2 if ord(char) > 127 else 1 for char in text)


def estimate_height(item: dict[str, Any]) -> float:
    name = str(item.get("name", ""))
    goal = str(item.get("goal", ""))
    progress = str(item.get("progress", ""))
    deviation = str(item.get("deviation", ""))
    lines = max(
        1,
        math.ceil(text_units(name) / 20),
        math.ceil(text_units(goal) / 32),
        math.ceil(text_units(progress) / 64),
        math.ceil(text_units(deviation) / 22),
    )
    return float(min(409, max(30, lines * 15 + 10)))


def vertical_label(label: str) -> str:
    compact = label.strip()
    if 1 < len(compact) <= 6 and all(
        "\u4e00" <= char <= "\u9fff" or char in " ·/-" for char in compact
    ):
        return "\n".join(compact)
    return compact


def apply_title(ws: Worksheet, period: str, title_prefix: str) -> None:
    ws.merge_cells("A1:F1")
    cell = ws["A1"]
    cell.value = f"{title_prefix}{period}"
    cell.font = Font(name=FONT_NAME, size=14, bold=True, color=TITLE_COLOR)
    cell.fill = TITLE_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    for column in range(1, 7):
        target = ws.cell(row=1, column=column)
        target.fill = TITLE_FILL
        target.border = BORDER
    ws.row_dimensions[1].height = 30


def apply_header(ws: Worksheet) -> None:
    for column, value in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=column, value=value)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = BORDER
    ws.row_dimensions[2].height = 30


def apply_section(
    ws: Worksheet,
    *,
    start_row: int,
    label: str,
    items: list[dict[str, Any]],
) -> int:
    end_row = start_row + len(items) - 1

    if end_row > start_row:
        ws.merge_cells(
            start_row=start_row,
            start_column=1,
            end_row=end_row,
            end_column=1,
        )

    label_cell = ws.cell(row=start_row, column=1, value=vertical_label(label))
    label_cell.font = Font(name=FONT_NAME, size=11, bold=True, color=TITLE_COLOR)
    label_cell.fill = SECTION_FILL
    label_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    for row in range(start_row, end_row + 1):
        target = ws.cell(row=row, column=1)
        target.fill = SECTION_FILL
        target.border = BORDER

    for offset, item in enumerate(items):
        row = start_row + offset
        values = [
            "",
            str(item.get("name", "")),
            str(item.get("goal", "")),
            str(item.get("progress", "")),
            str(item.get("status", "")),
            str(item.get("deviation", "")),
        ]
        for column, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=column)
            if column > 1:
                cell.value = value
            cell.font = Font(
                name=FONT_NAME,
                size=11,
                bold=(column == 2),
                color=TEXT_COLOR,
            )
            cell.alignment = Alignment(
                horizontal="center" if column == 5 else "left",
                vertical="center" if column == 1 else "top",
                wrap_text=True,
            )
            cell.border = BORDER
        ws.row_dimensions[row].height = estimate_height(item)

    return end_row


def build_workbook(report: dict[str, Any], *, sheet_name: str, title_prefix: str) -> tuple[Workbook, int]:
    validate_report(report)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    apply_title(worksheet, str(report["period"]), title_prefix)
    apply_header(worksheet)

    row = 3
    sections = report["sections"]
    for section_index, section in enumerate(sections):
        if section_index > 0:
            worksheet.row_dimensions[row].height = 8
            row += 1
        start_row = row
        end_row = apply_section(
            worksheet,
            start_row=start_row,
            label=str(section["label"]),
            items=section["items"],
        )
        row = end_row + 1

    for column, width in COLUMN_WIDTHS.items():
        worksheet.column_dimensions[column].width = width

    worksheet.sheet_view.showGridLines = False
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:2"
    worksheet.page_margins.left = 0.3
    worksheet.page_margins.right = 0.3
    worksheet.page_margins.top = 0.4
    worksheet.page_margins.bottom = 0.4

    return workbook, row - 1


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Render report.json into the fixed weekly-report Excel template.",
    )
    parser.add_argument("--input", required=True, help="report.json path")
    parser.add_argument(
        "--output",
        help="output .xlsx path; defaults to config output_dir + 周报_<period>.xlsx",
    )
    parser.add_argument(
        "--config",
        help=f"local config JSON (default: {DEFAULT_CONFIG})",
    )
    parser.add_argument(
        "--sheet-name",
        default=None,
        help="worksheet name; defaults to config report.sheet_name or 上周周报",
    )
    parser.add_argument(
        "--title-prefix",
        default="周报周期：",
        help="title prefix (default: 周报周期：)",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        config_path = Path(args.config).expanduser() if args.config else DEFAULT_CONFIG
        config = load_config(config_path)
        report = load_report(Path(args.input).expanduser())

        sheet_name = args.sheet_name
        if not sheet_name:
            report_config = config.get("report", {})
            if isinstance(report_config, dict):
                sheet_name = str(report_config.get("sheet_name", "")).strip() or None
        sheet_name = sheet_name or "上周周报"

        if args.output:
            output = Path(args.output).expanduser()
        else:
            output_dir = config.get("output_dir")
            if not output_dir:
                raise ReportError(
                    "--output is required when config has no output_dir"
                )
            output = Path(str(output_dir)).expanduser() / report_filename(
                str(report.get("period", ""))
            )

        workbook, last_row = build_workbook(
            report,
            sheet_name=sheet_name,
            title_prefix=args.title_prefix,
        )
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    except ReportError as exc:
        print(f"[error] {exc}", file=sys.stderr)
        return 2

    item_count = sum(len(section["items"]) for section in report["sections"])
    print(f"Saved: {output}")
    print(f"Sections: {len(report['sections'])}")
    print(f"Items: {item_count}")
    print(f"Rows: {last_row}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
