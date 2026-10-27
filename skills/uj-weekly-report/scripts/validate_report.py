#!/usr/bin/env python3
"""Validate a UJoin personal weekly report before delivery.

Three layers:

1. report.json: fields, fixed sections, status vocabulary, evidence.
2. coverage: every project with commits is referenced by at least one item.
3. Excel: the rendered file must match ``assets/ujoin-weekly-template.xlsx``
   cell by cell for styles, merges, widths, heights and data validation,
   and the data cells must match report.json.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    print(
        "[error] openpyxl is required. Run with: uv run --with openpyxl python ...",
        file=sys.stderr,
    )
    raise SystemExit(2)

from report_layout import (
    DEFAULT_CONFIG,
    DEFAULT_TEMPLATE,
    HEADERS,
    INSTRUCTIONS,
    LayoutError,
    PERIOD_PREFIX,
    PERIOD_ROW,
    SECOND_BLOCK,
    SECTIONS,
    TITLE_PREFIX,
    TITLE_ROW,
    load_template,
    normalize_status,
    render_deviation,
    resolve_person_name,
    section_map,
    text,
)


class Validation:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.warnings: list[str] = []

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warning(self, message: str) -> None:
        self.warnings.append(message)


def load_json(path: Path, label: str) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"{label} not found: {path}")
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid JSON in {label}: {path}: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError(f"{label} must be a JSON object: {path}")
    return data


def load_config(path: Path, *, required: bool) -> dict[str, Any]:
    if not path.exists():
        if required:
            raise FileNotFoundError(f"config not found: {path}")
        return {}
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"invalid config: {path}: {exc}") from exc
    if not isinstance(config, dict):
        raise ValueError(f"config must be a JSON object: {path}")
    return config


def check_report_structure(
    report: dict[str, Any],
    commits: dict[str, Any],
    result: Validation,
) -> tuple[set[str], int, dict[str, int]]:
    period = commits.get("period", {})
    report_period = text(report.get("period"))
    if not report_period:
        result.error("report.period is empty")
    elif report_period != text(period.get("label")):
        result.error(
            f"report.period {report_period!r} does not match commits.period.label "
            f"{text(period.get('label'))!r}"
        )

    report_week = text(report.get("week"))
    commits_week = text(period.get("week"))
    if report_week and commits_week and report_week != commits_week:
        result.error(
            f"report.week {report_week!r} does not match commits.period.week {commits_week!r}"
        )

    if period.get("multi_week"):
        result.error("the selected range spans more than one ISO week; use a single week")

    try:
        mapped = section_map(report)
    except LayoutError as exc:
        result.error(str(exc))
        return set(), 0, {}

    by_project = commits.get("stats", {}).get("by_project", {})
    if not isinstance(by_project, dict):
        by_project = {}

    covered_projects: set[str] = set()
    section_counts: dict[str, int] = {}
    total_items = 0

    for entry in SECTIONS:
        label = entry["label"]
        items = mapped[label]
        section_counts[label] = len(items)
        total_items += len(items)

        for index, item in enumerate(items, start=1):
            prefix = f"{label}[{index}]"
            for field in ("name", "goal", "progress", "status"):
                if not text(item.get(field)):
                    result.error(f"{prefix}.{field} is empty")

            status = ""
            try:
                status = normalize_status(item.get("status"))
            except LayoutError as exc:
                result.error(f"{prefix}.status: {exc}")

            deviation = render_deviation(item.get("deviation"))
            if status and status != "✅已完成" and not deviation:
                result.error(
                    f"{prefix}.deviation must explain why status is {status!r}"
                )
            if status == "✅已完成" and deviation:
                result.warning(
                    f"{prefix} is 已完成 but deviation is not empty: {deviation!r}"
                )

            evidence = item.get("evidence")
            if not isinstance(evidence, dict):
                result.error(f"{prefix}.evidence must be an object")
                continue

            projects = evidence.get("projects")
            if not isinstance(projects, list) or not projects:
                result.error(f"{prefix}.evidence.projects must be a non-empty list")
                projects = []
            normalized_projects = [text(project) for project in projects if text(project)]
            if len(normalized_projects) != len(projects):
                result.error(f"{prefix}.evidence.projects contains an empty project name")

            for project in normalized_projects:
                if project not in by_project:
                    result.error(
                        f"{prefix}.evidence.projects references {project!r}, "
                        "which has no commits in commits.json"
                    )
                else:
                    covered_projects.add(project)

            commit_count = evidence.get("commit_count")
            if not isinstance(commit_count, int) or isinstance(commit_count, bool):
                result.error(f"{prefix}.evidence.commit_count must be an integer")
            elif commit_count <= 0:
                result.error(f"{prefix}.evidence.commit_count must be greater than 0")
            else:
                actual = sum(int(by_project.get(project, 0)) for project in normalized_projects)
                if commit_count > actual:
                    result.error(
                        f"{prefix}.evidence.commit_count {commit_count} exceeds "
                        f"the actual {actual} commit(s) in {normalized_projects}"
                    )

            samples = evidence.get("sample_commits")
            if not isinstance(samples, list) or not samples:
                result.error(f"{prefix}.evidence.sample_commits must be a non-empty list")
            else:
                min_samples = (
                    1
                    if isinstance(commit_count, int)
                    and not isinstance(commit_count, bool)
                    and commit_count == 1
                    else 2
                )
                if not min_samples <= len(samples) <= 5:
                    result.warning(
                        f"{prefix}.evidence.sample_commits should have "
                        f"{min_samples}-5 items, got {len(samples)}"
                    )
                if not all(text(sample) for sample in samples):
                    result.error(f"{prefix}.evidence.sample_commits contains an empty value")

            repos = evidence.get("repos")
            if repos is not None:
                if not isinstance(repos, list) or not all(text(repo) for repo in repos):
                    result.error(f"{prefix}.evidence.repos must be a list of repo paths")
                else:
                    by_repo = commits.get("stats", {}).get("by_repo", {})
                    if isinstance(by_repo, dict):
                        for repo in repos:
                            repo_name = text(repo)
                            if repo_name and repo_name not in by_repo:
                                result.warning(
                                    f"{prefix}.evidence.repos references {repo_name!r}, "
                                    "which has no commits in commits.json"
                                )

    return covered_projects, total_items, section_counts


def check_coverage(
    commits: dict[str, Any],
    covered_projects: set[str],
    result: Validation,
) -> None:
    unmapped = commits.get("unmapped", [])
    if unmapped:
        repos = ", ".join(
            f"{item.get('repo')}({item.get('commit_count')})"
            for item in unmapped
            if isinstance(item, dict)
        )
        result.error(f"unmapped repos must be confirmed before generating the report: {repos}")

    ignored = commits.get("ignored", [])
    if isinstance(ignored, list) and ignored:
        ignored_count = sum(
            int(item.get("commit_count", 0))
            for item in ignored
            if isinstance(item, dict)
        )
        result.warning(
            f"ignored_repos excludes {ignored_count} own commit(s); "
            "confirm they are not reportable work"
        )

    stats = commits.get("stats", {})
    by_project = stats.get("by_project", {})
    if isinstance(by_project, dict):
        missing = sorted(
            project
            for project, count in by_project.items()
            if int(count) > 0 and project not in covered_projects
        )
        if missing:
            result.error(
                "projects with commits are not covered by any report item: "
                + ", ".join(missing)
            )

    if stats.get("errors"):
        for error in stats["errors"]:
            result.warning(f"extract warning: {error}")

    if stats.get("missing_repos"):
        result.warning(
            "project map contains repos that do not exist in the workspace: "
            + ", ".join(stats["missing_repos"])
        )

    total = stats.get("total")
    commit_list = commits.get("commits", [])
    if isinstance(total, int) and isinstance(commit_list, list) and total != len(commit_list):
        result.warning(f"stats.total {total} does not match len(commits) {len(commit_list)}")


def check_recommended_counts(
    section_counts: dict[str, int],
    total_items: int,
    result: Validation,
) -> None:
    key_count = section_counts.get("重点项目")
    daily_count = section_counts.get("日常工作")
    if key_count is not None and not 3 <= key_count <= 4:
        result.warning(f"重点项目 should have 3-4 items, got {key_count}")
    if daily_count is not None and not 2 <= daily_count <= 4:
        result.warning(f"日常工作 should have 2-4 items, got {daily_count}")
    if total_items > 8:
        result.warning(f"total items must be <= 8 (4+4 template slots), got {total_items}")


def check_samples(
    report: dict[str, Any],
    commits: dict[str, Any],
    result: Validation,
) -> None:
    by_project: dict[str, set[str]] = {}
    for commit in commits.get("commits", []):
        if not isinstance(commit, dict):
            continue
        project = text(commit.get("project"))
        if not project:
            continue
        by_project.setdefault(project, set()).add(text(commit.get("norm")))
        by_project[project].add(text(commit.get("subject")))

    for section in report.get("sections", []):
        if not isinstance(section, dict):
            continue
        for item in section.get("items", []):
            if not isinstance(item, dict):
                continue
            evidence = item.get("evidence", {})
            projects = evidence.get("projects", []) if isinstance(evidence, dict) else []
            samples = evidence.get("sample_commits", []) if isinstance(evidence, dict) else []
            corpus: set[str] = set()
            for project in projects:
                corpus.update(by_project.get(text(project), set()))
            if not corpus or not isinstance(samples, list):
                continue
            matched = 0
            for sample in samples:
                sample_text = text(sample)
                if not sample_text:
                    continue
                if any(sample_text in entry or entry in sample_text for entry in corpus if entry):
                    matched += 1
            if matched == 0:
                result.warning(
                    f"sample_commits for {text(item.get('name'))!r} do not match any commit title"
                )


def color_signature(color: Any) -> Any:
    if color is None:
        return None
    return (
        color.type,
        color.rgb if color.type == "rgb" else None,
        color.theme if color.type == "theme" else None,
        round(color.tint, 6) if color.type == "theme" else None,
        color.indexed if color.type == "indexed" else None,
    )


def style_signature(cell: Any) -> Any:
    font = cell.font
    fill = cell.fill
    border = cell.border
    align = cell.alignment
    return (
        font.name,
        font.sz,
        font.b,
        font.i,
        font.u,
        color_signature(font.color),
        fill.patternType,
        color_signature(fill.fgColor),
        color_signature(fill.bgColor),
        border.left.style,
        border.right.style,
        border.top.style,
        border.bottom.style,
        color_signature(border.left.color),
        color_signature(border.right.color),
        color_signature(border.top.color),
        color_signature(border.bottom.color),
        align.horizontal,
        align.vertical,
        align.wrapText,
        align.textRotation,
        align.indent,
        cell.number_format,
    )


def check_template_fidelity(
    worksheet: Any,
    template_ws: Any,
    result: Validation,
) -> None:
    expected_merges = sorted(str(item) for item in template_ws.merged_cells.ranges)
    actual_merges = sorted(str(item) for item in worksheet.merged_cells.ranges)
    if actual_merges != expected_merges:
        result.error(
            f"merged ranges do not match the template: {actual_merges!r}"
        )

    for column in "ABCDEF":
        expected = template_ws.column_dimensions[column].width
        actual = worksheet.column_dimensions[column].width
        if expected is None and actual is None:
            continue
        if expected is None or actual is None or abs(float(actual) - float(expected)) > 0.01:
            result.error(
                f"column {column} width should be {expected!r} like the template, "
                f"got {actual!r}"
            )

    for row in range(1, 23):
        expected = template_ws.row_dimensions[row].height
        actual = worksheet.row_dimensions[row].height
        if expected != actual:
            result.error(
                f"row {row} height should be {expected!r} like the template, got {actual!r}"
            )

    for row in range(1, 23):
        for column in range(1, 7):
            expected_cell = template_ws.cell(row=row, column=column)
            actual_cell = worksheet.cell(row=row, column=column)
            if style_signature(expected_cell) != style_signature(actual_cell):
                result.error(
                    f"cell {actual_cell.coordinate} style does not match the template"
                )

    if worksheet.sheet_view.showGridLines != template_ws.sheet_view.showGridLines:
        result.error(
            f"gridlines should be {template_ws.sheet_view.showGridLines}, "
            f"got {worksheet.sheet_view.showGridLines}"
        )

    if worksheet.page_setup.orientation != template_ws.page_setup.orientation:
        result.error("page orientation must stay unset like the template")
    if worksheet.sheet_properties.pageSetUpPr and worksheet.sheet_properties.pageSetUpPr.fitToPage:
        result.error("fitToPage must stay unset like the template")

    expected_validations = sorted(
        (str(item.sqref), text(item.formula1))
        for item in template_ws.data_validations.dataValidation
    )
    actual_validations = sorted(
        (str(item.sqref), text(item.formula1))
        for item in worksheet.data_validations.dataValidation
    )
    if actual_validations != expected_validations:
        result.error(
            "data validation does not match the template: "
            f"{actual_validations!r}"
        )


def check_excel(
    xlsx_path: Path,
    report: dict[str, Any],
    config: dict[str, Any],
    template_path: Path,
    result: Validation,
    *,
    person_name: str | None = None,
) -> None:
    if not xlsx_path.exists():
        result.error(f"xlsx not found: {xlsx_path}")
        return

    workbook = load_workbook(xlsx_path)
    name = text(person_name) or resolve_person_name(report, config)
    if len(workbook.sheetnames) != 1:
        result.error(
            f"workbook should contain exactly one sheet, got {workbook.sheetnames!r}"
        )
    sheet_name = workbook.sheetnames[0]
    worksheet = workbook[sheet_name]
    if sheet_name != name:
        result.warning(f"sheet name {sheet_name!r} does not match person {name!r}")

    expected_title = f"{TITLE_PREFIX}{name}"
    if text(worksheet.cell(row=TITLE_ROW, column=1).value) != expected_title:
        result.error(
            f"A1 should be {expected_title!r}, "
            f"got {text(worksheet.cell(row=TITLE_ROW, column=1).value)!r}"
        )

    if text(worksheet.cell(row=2, column=1).value) != INSTRUCTIONS:
        result.error("A2 instructions do not match the template")

    expected_period = f"{PERIOD_PREFIX}{text(report.get('period'))}"
    if text(worksheet.cell(row=PERIOD_ROW, column=1).value) != expected_period:
        result.error(
            f"A{PERIOD_ROW} should be {expected_period!r}, "
            f"got {text(worksheet.cell(row=PERIOD_ROW, column=1).value)!r}"
        )

    headers = [
        text(worksheet.cell(row=4, column=column).value) for column in range(1, 7)
    ]
    if headers != HEADERS:
        result.error(f"header row does not match the template: {headers!r}")

    try:
        mapped = section_map(report)
    except LayoutError as exc:
        result.error(str(exc))
        return

    for entry in SECTIONS:
        items = mapped[entry["label"]]
        for offset in range(entry["slots"]):
            row = entry["start_row"] + offset
            item = items[offset] if offset < len(items) else None
            if item is None:
                for column in range(2, 7):
                    if worksheet.cell(row=row, column=column).value is not None:
                        result.error(
                            f"{worksheet.cell(row=row, column=column).coordinate} "
                            "should be empty to fill the template slot"
                        )
                continue
            expected = {
                2: text(item.get("name")),
                3: text(item.get("goal")),
                4: text(item.get("progress")),
                5: normalize_status(item.get("status")),
                6: render_deviation(item.get("deviation")),
            }
            for column, expected_value in expected.items():
                actual_value = text(worksheet.cell(row=row, column=column).value)
                if actual_value != expected_value:
                    result.error(
                        f"{worksheet.cell(row=row, column=column).coordinate} "
                        f"should be {expected_value!r}, got {actual_value!r}"
                    )

    if text(worksheet.cell(row=SECOND_BLOCK["period_row"], column=1).value) != (
        f"{PERIOD_PREFIX}2026/月/日-2026/月/日"
    ):
        result.error("the lower blank table must keep the template period placeholder")
    lower_headers = [
        text(worksheet.cell(row=SECOND_BLOCK["header_row"], column=column).value)
        for column in range(1, 7)
    ]
    if lower_headers != HEADERS:
        result.error("the lower blank table headers were modified")
    for row in range(SECOND_BLOCK["start_row"], SECOND_BLOCK["end_row"] + 1):
        for column in range(2, 7):
            if worksheet.cell(row=row, column=column).value is not None:
                result.error(
                    f"{worksheet.cell(row=row, column=column).coordinate} "
                    "in the lower blank table must stay empty"
                )

    template = load_template(template_path)
    template_ws = template[template.sheetnames[0]]
    check_template_fidelity(worksheet, template_ws, result)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Validate report.json, commits.json and the rendered UJoin weekly report.",
    )
    parser.add_argument("--report", required=True, help="report.json path")
    parser.add_argument("--commits", required=True, help="commits.json path")
    parser.add_argument("--xlsx", help="rendered .xlsx path; omit to skip Excel checks")
    parser.add_argument(
        "--config",
        help=f"local config JSON (default: {DEFAULT_CONFIG})",
    )
    parser.add_argument(
        "--template",
        help=f"template .xlsx (default: {DEFAULT_TEMPLATE})",
    )
    parser.add_argument("--person-name", help="override the name in the title")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    result = Validation()

    try:
        config_path = Path(args.config).expanduser() if args.config else DEFAULT_CONFIG
        config = load_config(config_path, required=bool(args.config))
        report = load_json(Path(args.report).expanduser(), "report")
        commits = load_json(Path(args.commits).expanduser(), "commits")
    except (FileNotFoundError, ValueError) as exc:
        print(f"[error] {exc}", file=sys.stderr)
        return 2

    template_value = args.template or (
        (config.get("report") or {}).get("template")
        if isinstance(config.get("report"), dict)
        else None
    )
    template_path = (
        Path(str(template_value)).expanduser() if template_value else DEFAULT_TEMPLATE
    )

    covered_projects, total_items, section_counts = check_report_structure(
        report, commits, result
    )
    check_coverage(commits, covered_projects, result)
    check_recommended_counts(section_counts, total_items, result)
    check_samples(report, commits, result)

    if args.xlsx:
        check_excel(
            Path(args.xlsx).expanduser(),
            report,
            config,
            template_path,
            result,
            person_name=args.person_name,
        )
    else:
        result.warning("xlsx not provided; Excel layout was not checked")

    if result.errors:
        print(f"[FAIL] {len(result.errors)} error(s), {len(result.warnings)} warning(s)")
        for error in result.errors:
            print(f"  error: {error}")
        for warning in result.warnings:
            print(f"  warning: {warning}")
        return 1

    print(f"[PASS] {len(result.warnings)} warning(s)")
    for warning in result.warnings:
        print(f"  warning: {warning}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
