#!/usr/bin/env python3
"""Shared layout contract for the UJoin personal weekly report template.

The source of truth is ``assets/ujoin-weekly-template.xlsx``, which is a
sanitized copy of ``ujoin周报模板.xlsx``. This module only describes the
fixed slots; it does not touch any cell.

Template structure (sheet ``个人周报（模板）``)::

    row 1  个人周报——员工姓名        merged A1:F1
    row 2  填写说明                   merged A2:F2
    row 3  周报周期：...              merged A3:F3
    row 4  大类 / 名称（只列关键） / 目标 / 进展 / 完成状态 / 偏差说明
    row 5-8   重点项目（4 个固定槽位，A5:A8 合并）
    row 9-12  日常工作（4 个固定槽位，A9:A12 合并）
    row 13-22 第二张空白表，保持模板原样
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "ujoin-weekly-template.xlsx"
DEFAULT_CONFIG = Path.home() / ".config" / "git-weekly-report" / "config.json"

BASE_SHEET = "个人周报（模板）"
TITLE_PREFIX = "个人周报——"
PERIOD_PREFIX = "周报周期："
DEFAULT_PERSON = "员工姓名"
INSTRUCTIONS = (
    "1. 本表是模板，不要直接填写，请创建副本后使用。\n"
    "2. 最晚每周日18点前更新完，如本部门周会提前或有其他要求，以主管要求为准。"
)

TITLE_ROW = 1
INSTRUCTIONS_ROW = 2
PERIOD_ROW = 3
HEADER_ROW = 4

HEADERS = ["大类", "名称\n（只列关键）", "目标", "进展", "完成状态", "偏差说明"]
COLUMN_WIDTHS = {"B": 23.0625, "C": 37.5312, "D": 61.3203}

SECTIONS = [
    {"label": "重点项目", "start_row": 5, "slots": 4},
    {"label": "日常工作", "start_row": 9, "slots": 4},
]
SECOND_BLOCK = {"period_row": 13, "header_row": 14, "start_row": 15, "end_row": 22}
LAST_ROW = 22

STATUS_DONE = "✅已完成"
STATUS_DOING = "🟡进行中"
STATUS_PAUSED = "⏸暂停"
STATUS_DELAYED = "❌延期"
STATUS_VALUES = [STATUS_DONE, STATUS_DOING, STATUS_PAUSED, STATUS_DELAYED]

STATUS_ALIASES = {
    "完成": STATUS_DONE,
    "已完成": STATUS_DONE,
    STATUS_DONE: STATUS_DONE,
    "进行中": STATUS_DOING,
    STATUS_DOING: STATUS_DOING,
    "部分完成": STATUS_DOING,
    "暂停": STATUS_PAUSED,
    STATUS_PAUSED: STATUS_PAUSED,
    "延期": STATUS_DELAYED,
    "有偏差": STATUS_DELAYED,
    STATUS_DELAYED: STATUS_DELAYED,
}

BLANK_DEVIATIONS = {"", "无", "无。", "-", "none", "n/a"}

MERGES = [
    "A1:F1",
    "A2:F2",
    "A3:F3",
    "A5:A8",
    "A9:A12",
    "A13:F13",
    "A15:A18",
    "A19:A22",
]
VALIDATION_RANGES = ["E5:E12", "E15:E22"]


class LayoutError(RuntimeError):
    """Raised when report data cannot fit the fixed template."""


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_status(value: Any) -> str:
    raw = text(value)
    if raw in STATUS_ALIASES:
        return STATUS_ALIASES[raw]
    raise LayoutError(
        f"unsupported status {raw!r}; use one of {', '.join(STATUS_VALUES)}"
    )


def render_deviation(value: Any) -> str:
    raw = text(value)
    return "" if raw.lower() in BLANK_DEVIATIONS else raw


def resolve_person_name(report: dict[str, Any], config: dict[str, Any]) -> str:
    for candidate in (
        report.get("person_name"),
        (config.get("report") or {}).get("person_name")
        if isinstance(config.get("report"), dict)
        else None,
    ):
        name = text(candidate)
        if name:
            return name
    return DEFAULT_PERSON


def report_filename(person_name: str, period: str) -> str:
    import re

    match = re.fullmatch(
        r"(\d{4})/(\d{2})/(\d{2})-(\d{4})/(\d{2})/(\d{2})",
        period.strip(),
    )
    if not match:
        raise LayoutError(
            f"cannot derive output filename from period {period!r}; pass --output"
        )
    start = "".join(match.groups()[:3])
    end = "".join(match.groups()[3:])
    safe_name = re.sub(r"[\\/:*?\"<>|]", "_", person_name)
    return f"个人周报_{safe_name}_{start}-{end}.xlsx"


def load_template(path: Path) -> Any:
    """Load the .xlsx template, patching the WPS ``<fill/>`` quirk if needed."""
    from openpyxl import load_workbook

    if not path.exists():
        raise LayoutError(f"template not found: {path}")
    try:
        return load_workbook(path)
    except TypeError:
        patched = patch_wps_styles(path)
        try:
            return load_workbook(patched)
        finally:
            patched.unlink(missing_ok=True)


def patch_wps_styles(path: Path) -> Path:
    """WPS writes ``<fill/>`` which openpyxl rejects. Fix a throwaway copy."""
    import os
    import tempfile
    import zipfile

    handle, temp_name = tempfile.mkstemp(suffix=".xlsx")
    os.close(handle)
    target = Path(temp_name)
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w") as out:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = data.replace(
                    b"<fill/>", b'<fill><patternFill patternType="none"/></fill>'
                ).replace(
                    b"<fill></fill>",
                    b'<fill><patternFill patternType="none"/></fill>',
                )
            out.writestr(item, data)
    return target


def section_map(report: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    sections = report.get("sections")
    if not isinstance(sections, list) or not sections:
        raise LayoutError("report.sections must be a non-empty list")

    mapped: dict[str, list[dict[str, Any]]] = {}
    for index, section in enumerate(sections, start=1):
        if not isinstance(section, dict):
            raise LayoutError(f"sections[{index}] must be an object")
        label = text(section.get("label"))
        items = section.get("items")
        if not isinstance(items, list):
            raise LayoutError(f"sections[{index}].items must be a list")
        mapped[label] = items

    expected = [entry["label"] for entry in SECTIONS]
    missing = [label for label in expected if label not in mapped]
    extra = [label for label in mapped if label not in expected]
    if missing or extra:
        raise LayoutError(
            "sections must be exactly "
            + " / ".join(expected)
            + f"; missing={missing or 'none'}, extra={extra or 'none'}"
        )

    for entry in SECTIONS:
        items = mapped[entry["label"]]
        if len(items) > entry["slots"]:
            raise LayoutError(
                f"{entry['label']} has {len(items)} items but the template only "
                f"has {entry['slots']} rows; merge items before rendering"
            )
    return mapped
